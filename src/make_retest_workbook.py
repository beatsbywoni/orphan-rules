# -*- coding: utf-8 -*-
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SRC = "goldset_retest_sept_blank.csv"
OUT = "goldset_retest_sept.xlsx"
FONT = "Arial"

rows = list(csv.DictReader(open(SRC, encoding="utf-8-sig")))
cols = list(rows[0].keys())
assert cols[-4:] == ["L1_위임인가", "L2_위임된법형식", "L3_확신도", "L4_메모"], cols[-4:]

L1 = ["위임", "인용", "판단불가"]
L2 = ["대통령령", "부령", "고시", "훈령", "예규", "조례", "교육규칙", "미특정", "해당없음",
      "대통령령|조례", "대통령령|교육규칙", "대통령령|부령", "조례|교육규칙"]

wb = Workbook()

# ─────────────────────────── 안내 ───────────────────────────
g = wb.active
g.title = "안내"
YEL = PatternFill("solid", fgColor="FFF2CC")
HDR = Font(name=FONT, bold=True, size=13)
BODY = Font(name=FONT, size=11)

lines = [
    ("9월 재검사 — 읽고 시작하십시오", "h"),
    ("", ""),
    ("이 파일은 사전등록된 intra-annotator 재검사용입니다. 2026-08-15에 라벨한 200건 중", "b"),
    ("무작위 60건을 3주 이상 간격을 두고 다시 라벨합니다. 8월 재검사와 겹치는 항목은 0건입니다.", "b"),
    ("", ""),
    ("규칙 세 가지", "h2"),
    ("1. 이전 라벨을 절대 보지 마십시오. data/goldset_labeled_v2.csv 를 열지 마십시오.", "b"),
    ("2. 기억이 나더라도 기억으로 답하지 말고, 코드북 시트의 기준으로 다시 판정하십시오.", "b"),
    ("3. 한 번에 끝내십시오. 중간에 이전 결과를 확인하고 돌아오면 측정이 무의미해집니다.", "b"),
    ("", ""),
    ("채울 곳", "h2"),
    ("[라벨링] 시트의 노란색 네 열만 채웁니다. 나머지 열은 건드리지 마십시오.", "b"),
    ("  · L1_위임인가   드롭다운: 위임 / 인용 / 판단불가", "b"),
    ("  · L2_위임된법형식  드롭다운. 두 형식이 병기되면 | 로 이어 씁니다 (예: 대통령령|조례).", "b"),
    ("                  목록에 없는 조합도 직접 입력할 수 있습니다.", "b"),
    ("  · L3_확신도     1(전혀 확신 없음) ~ 5(확실). 3 이하는 조정 대상입니다.", "b"),
    ("  · L4_메모       판단이 갈릴 만한 이유를 한 줄. 비워도 됩니다.", "b"),
    ("", ""),
    ("작성 예시는 [코드북 v1.1] 시트 아래쪽 '고정 예시 4'에 있습니다.", "b"),
    ("이 예시는 8월 라벨링에 쓰인 것과 같은 것으로, 새로운 정보가 아닙니다.", "b"),
    ("", ""),
    ("끝나면", "h2"),
    ("파일을 그대로 저장한 뒤, 프로젝트 폴더에서 한 줄만 실행하십시오.", "b"),
    ("    sh run_sept_retest.sh", "m"),
    ("κ 산출과 원고 문장 생성까지 자동으로 됩니다.", "b"),
    ("", ""),
    ("진행 상황", "h2"),
]
r = 1
for text, kind in lines:
    c = g.cell(row=r, column=1, value=text)
    c.font = HDR if kind == "h" else (Font(name=FONT, bold=True, size=11) if kind == "h2"
             else (Font(name="Consolas", size=11) if kind == "m" else BODY))
    r += 1
g.cell(row=r, column=1, value="L1 입력 완료").font = BODY
g.cell(row=r, column=2, value="=COUNTA(라벨링!K2:K61)").font = BODY
g.cell(row=r + 1, column=1, value="L2 입력 완료").font = BODY
g.cell(row=r + 1, column=2, value="=COUNTA(라벨링!L2:L61)").font = BODY
g.cell(row=r + 2, column=1, value="남은 건수 (L1)").font = BODY
g.cell(row=r + 2, column=2, value="=60-COUNTA(라벨링!K2:K61)").font = Font(name=FONT, bold=True, size=11)
g.column_dimensions["A"].width = 92
g.column_dimensions["B"].width = 12

# ─────────────────────────── 라벨링 ───────────────────────────
ws = wb.create_sheet("라벨링")
thin = Side(style="thin", color="BFBFBF")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
HFILL = PatternFill("solid", fgColor="D9D9D9")

for j, name in enumerate(cols, start=1):
    c = ws.cell(row=1, column=j, value=name)
    c.font = Font(name=FONT, bold=True, size=10)
    c.fill = HFILL
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = BORD

for i, rec in enumerate(rows, start=2):
    for j, name in enumerate(cols, start=1):
        c = ws.cell(row=i, column=j, value=rec[name] or None)
        c.font = Font(name=FONT, size=10)
        c.border = BORD
        c.alignment = Alignment(vertical="top",
                                wrap_text=(name in ("조문문장", "상위조문제목", "대상명", "L4_메모")))
        if name in ("L1_위임인가", "L2_위임된법형식", "L3_확신도", "L4_메모"):
            c.fill = YEL

widths = {"id": 8, "상위법령명": 22, "상위조문번호": 7, "상위조문제목": 20, "조문문장": 68,
          "위임표지어구": 14, "API_위임구분": 11, "자동추출_위임형식": 12, "대상유형": 9,
          "대상명": 26, "L1_위임인가": 12, "L2_위임된법형식": 15, "L3_확신도": 9, "L4_메모": 26}
for j, name in enumerate(cols, start=1):
    ws.column_dimensions[get_column_letter(j)].width = widths.get(name, 14)
ws.row_dimensions[1].height = 30
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"

def dv(values, col, warn_title, warn_msg):
    d = DataValidation(type="list", formula1='"' + ",".join(values) + '"',
                       allow_blank=True, showDropDown=False)
    d.errorStyle = "warning"
    d.errorTitle = warn_title
    d.error = warn_msg
    ws.add_data_validation(d)
    d.add(f"{col}2:{col}{len(rows)+1}")

k = get_column_letter(cols.index("L1_위임인가") + 1)
l = get_column_letter(cols.index("L2_위임된법형식") + 1)
m = get_column_letter(cols.index("L3_확신도") + 1)
dv(L1, k, "목록에 없는 값", "L1은 위임 / 인용 / 판단불가 셋 중 하나입니다.")
dv(L2, l, "목록에 없는 값", "코드북 값 목록을 확인하십시오. 병기 형식은 | 로 잇습니다 (예: 대통령령|조례).")
dv(["1", "2", "3", "4", "5"], m, "범위 밖", "확신도는 1~5입니다.")

# ─────────────────────────── 코드북 ───────────────────────────
cb = wb.create_sheet("코드북 v1.1")
codebook = """코드북 v1.1 (2026-08-15 동결) — 8월 라벨링에 쓰인 것과 동일

[L1] 이 문장이 하위 규범에 규율권한을 넘기는가?
  위임      = 하위 규범이 정할 사항의 범위를 지시하며 규율권한을 넘긴다.
  인용      = 다른 법령을 가리키기만 한다. 정의 원용, 절차 준용, 「○○법」에 따른 형태.
  판단불가  = 문장이 잘렸거나 맥락 없이 판정할 수 없다.

  경계 규칙
  - 준용 규정("제○조를 준용한다")은 → 인용
  - "대통령령으로 정하는 ○○" 처럼 요건 일부만 맡기는 부분 위임 → 위임
  - 규범 제정이 아니라 개별 처분 권한만 부여하는 경우 → 인용
  - 포괄위임이 의심되어도 일단 위임. 위헌 여부는 판단하지 않는다.
  - 잘린 단편이라도 위임 표지어구가 있으면 그 표지로 판정한다 (절단 규칙).

[L2] 문장이 명시적으로 지정한 법형식
  실제로 무엇이 제정되었는지가 아니라, 문언이 무엇을 요구하는지로 판정한다.

  값 목록 (이 아홉 개가 전부다)
    대통령령 / 부령 / 고시 / 훈령 / 예규 / 조례 / 교육규칙 / 미특정 / 해당없음

  - 조례(지방의회 제정)와 교육규칙(교육감 제정)은 구분한다.
  - 복수 법형식이 병기된 경우 전부 | 로 연결한다.
      예: "대통령령 또는 조례로 정하는 바에 따라"      → 대통령령|조례
          "대통령령으로 정하는 범위에서 조례로 정한다"  → 대통령령|조례
  - 미특정: "장관이 정하는", "교육감이 지정하는" 등 형식을 지정하지 않은 경우.
            이것이 본 연구의 핵심 회색지대이므로 정확히 구분할 것.
  - 해당없음: L1이 인용인 경우.

[L3] 확신도 1~5. 3 이하는 조정 대상.
[L4] 판단이 갈릴 만한 이유를 한 줄.

────────────────────────────────────────────────────────────
고정 예시 4 (8월과 동일)

예시 1
  문장: 학교의 설립 기준에 관하여 필요한 사항은 대통령령으로 정한다.
  답:   L1=위임  L2=대통령령  L3=5
  이유: 문말 종결형으로 대통령령에 규율권한을 넘긴다.

예시 2
  문장: 교육부령으로 정한 기준을 충족하는 기관은 지원 대상이 된다.
  답:   L1=인용  L2=해당없음  L3=5
  이유: 완료형 '정한'은 이미 존재하는 규범을 가리키는 인용이다.

예시 3
  문장: 수업료는 대통령령으로 정하는 범위에서 조례로 정한다.
  답:   L1=위임  L2=대통령령|조례  L3=4
  이유: 두 법형식이 병기된 위임으로 | 로 연결한다.

예시 4
  문장: 제12조부터 제15조까지의 규정을 준용한다.
  답:   L1=인용  L2=해당없음  L3=5
  이유: 준용 규정은 인용이다."""

for i, line in enumerate(codebook.split("\n"), start=1):
    c = cb.cell(row=i, column=1, value=line)
    c.font = Font(name=FONT, size=11,
                  bold=line.startswith("[") or line.startswith("코드북") or line.startswith("고정"))
cb.column_dimensions["A"].width = 96

wb.save(OUT)
print("wrote", OUT, "| rows:", len(rows))
