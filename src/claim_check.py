#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""claim_check.py — 대외구속지향 고아규칙(D1∩T1) 중 교육부 법령을 근거로 주장하는 규칙의
주장 조문 원문 대조 (등록 외 사후 분석, §7.2 보강).

질문: "recorded delegation 이 없다"는 헤드라인이 (a) 정부 메타데이터 누락인지
(b) 법령 원문에도 위임이 없는 것인지, 주장 조문을 직접 읽어 가른다.

  python3 src/claim_check.py            # 워크시트 생성 data/claim_check_blank.csv / .xlsx
  python3 src/claim_check.py --analyse  # data/claim_check_labeled.csv 집계 → docs/39_claim_check_result.md

판정 코드 (코드북 v1.1 L1 기준을 '주장 조문'에 적용):
  a   조문(또는 같은 법령의 다른 조문)에 이 규칙의 규율 사항을 장관/행정규칙에 맡기는
      위임 문언이 있다 → 메타데이터 누락 (delegation present in text)
  b   위임 문언이 없다 → 근거 과장. 하위코드:
        b-decree     위임은 대통령령·부령에만 가 있고 행정규칙에는 없다 (형식 불일치)
        b-authorise  기준·지침을 '마련/통보'할 권한만 있고 법규 위임은 아니다 (2020두43722 유형)
        b-unrelated  주장 조문이 규칙의 규율 사항과 무관 (개정으로 번호가 밀린 낡은 주장 포함)
        b-deleted    주장 조문이 삭제됨
        b-cite       단순 인용·유보
  c   판단 불가 (법령이 코퍼스 밖·조문 특정 불가·별표에만 있음)
"""
import csv, json, re, sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
TRIAGE = DATA / "findings_D1_triage_v3.csv"
D4 = DATA / "findings_D4_falseclaim.csv"
BODIES = DATA / "lawbodies.jsonl"
RULES = DATA / "admrules.jsonl"
BLANK = DATA / "claim_check_blank.csv"
LABELED = DATA / "claim_check_labeled.csv"
DOC = ROOT / "docs/39_claim_check_result.md"

MARK = re.compile(r"(교육부장관이 정하|장관이 정하|정하여 고시|고시한다|고시할 수|정하는 바에 따라|정한다|정할 수 있다|위탁|위임)")
STOP = set("에 관한 규정 지침 요령 고시 기준 운영 및 등 의 에 따른 관리 지급 제 차 년 년도 규칙".split())


def norm(s): return re.sub(r"[\s·ㆍ()（）「」]", "", s or "")


def load():
    tri = list(csv.DictReader(open(TRIAGE, encoding="utf-8-sig")))
    ext = {r["행정규칙일련번호"]: r for r in tri
           if r["T1_규범성격"] == "대외구속지향" and r["근거주장범주_자동"] == "교육부법령_근거주장"}
    d4 = [r for r in csv.DictReader(open(D4, encoding="utf-8-sig")) if r["행정규칙일련번호"] in ext]
    laws = {}
    for line in open(BODIES, encoding="utf-8"):
        o = json.loads(line); laws[norm(o["법령명한글"])] = o
    rules = {}
    for line in open(RULES, encoding="utf-8"):
        o = json.loads(line); rules[str(o["행정규칙일련번호"])] = o
    return ext, d4, laws, rules


def art_text(a):
    return re.sub(r"\s+", " ", (a.get("조내용") or "")).strip()


def find_article(law, no):
    for a in law["조문"]:
        if a.get("조번호") == str(no) and not a.get("가지번호"):
            return a
    return None


def first_article(rule):
    t = rule.get("_조문텍스트") or ""
    if not t.strip():
        t = json.dumps(rule.get("_본문", ""), ensure_ascii=False)
    t = re.sub(r"<[^>]+>", " ", t)
    m = re.search(r"제1조.{0,600}?(?=제2조|$)", t, re.S)
    return re.sub(r"\s+", " ", m.group(0) if m else t[:600])


CIT = re.compile(r"[「『]\s*([^」』]{2,60}?)\s*[」』]\s*(?:\([^)]*\)|（[^）]*）)?\s*(시행령|시행규칙)?\s*((?:제?\s*\d+\s*조(?:의\s*\d+)?(?:\s*제\s*\d+\s*항)?(?:\s*(?:제\s*\d+\s*호)?)?\s*(?:,|및|ㆍ|·|와|과|\s)*)+)")
UNB = re.compile(r"(?<![「『가-힣])((?:[가-힣ㆍ·]+\s?){1,9}?(?:법률|법|규정|령|규칙))\s*(제\s*\d+\s*조(?:의\s*\d+)?(?:\s*제\s*\d+\s*항)?)")
ART = re.compile(r"제?\s*(\d+)\s*조(?:의\s*(\d+))?(?:\s*제\s*(\d+)\s*항)?")
ALIAS = {"학점인정법": "학점인정 등에 관한 법률"}


def parse_citations(first, laws):
    """규칙 제1조에서 법령 제N조(의M)(제K항) 인용을 전부 뽑아 코퍼스 조문으로 해소한다.
    「」 인용, 「」 뒤의 시행령/시행규칙, '같은 법/동 시행령·시행규칙', 괄호 없는 법령명 인용을 다룬다."""
    out, last_law = [], None
    for m in CIT.finditer(first):
        name = m.group(1).strip() + ((" " + m.group(2)) if m.group(2) else ""); last_law = m.group(1).strip()
        for a in ART.finditer(m.group(3)):
            out.append((name, a.group(1), a.group(2) or "", a.group(3) or ""))
    for m in re.finditer(r"(같은\s*법|동)\s*(시행령|시행규칙)\s*((?:제?\s*\d+\s*조(?:의\s*\d+)?(?:\s*제\s*\d+\s*항)?(?:,|및|ㆍ|·|\s)*)+)", first):
        if last_law:
            base = re.sub(r"\s*(시행령|시행규칙)$", "", last_law)
            for a in ART.finditer(m.group(3)):
                out.append((base + " " + m.group(2), a.group(1), a.group(2) or "", a.group(3) or ""))
    stripped = CIT.sub(" ", first)
    for m in UNB.finditer(stripped):
        name = m.group(1).strip()
        if name.startswith(("같은", "동 ")) or name in ("규정", "법"): continue
        toks = name.split()
        for k in range(len(toks)):            # 앞쪽 군더더기('본 예규는 …')를 떼며 코퍼스 법령명에 맞춘다
            cand = " ".join(toks[k:])
            if norm(cand) in laws: name = cand; break
        else:
            name = " ".join(t for t in toks if not t.endswith(("는", "은")) and t not in ("이", "본"))
        a = ART.search(m.group(2)); last_law = name
        out.append((name, a.group(1), a.group(2) or "", a.group(3) or ""))
    if "별표" in first:
        out.append(("(별표 인용 있음 — 원문은 별표를 봐야 함)", "", "", ""))
    rows, seen = [], set()
    for name, no, gaji, hang in out:
        key = (norm(name), no, gaji, hang)
        if key in seen: continue
        seen.add(key)
        if not no:
            rows.append((name, "별표", "")); continue
        law = laws.get(norm(name)) or (laws.get(norm(ALIAS[name])) if name in ALIAS else None)
        label = f"{name} 제{no}조" + (f"의{gaji}" if gaji else "") + (f"제{hang}항" if hang else "")
        if law is None:
            rows.append((label, "코퍼스 밖(타부처·비법령·명칭 불일치)", "")); continue
        art = next((a for a in law["조문"] if a.get("조번호") == no and (a.get("가지번호") or "") == gaji), None)
        if art is None or not art_text(art):
            rows.append((label, "현행 조문 없음(삭제·번호 변경?)", "")); continue
        rows.append((label, "확인", art_text(art)))
    return rows


def cited_laws(first, laws):
    names = set()
    for m in CIT.finditer(first):
        n = m.group(1).strip()
        for cand in (n, n + " 시행령", n + " 시행규칙", ALIAS.get(n, "")):
            if cand and norm(cand) in laws: names.add(norm(cand))
    for m in UNB.finditer(CIT.sub(" ", first)):
        if norm(m.group(1)) in laws: names.add(norm(m.group(1)))
    return [laws[n] for n in sorted(names)]


def candidates(law, rule_name, exclude_no=None, k=4):
    toks = [w for w in re.split(r"[\s·ㆍ()「」,]+", rule_name) if len(w) >= 2 and w not in STOP]
    scored = []
    for a in law["조문"]:
        if a.get("가지번호"): continue
        t = art_text(a)
        if not t: continue
        s = sum(1 for w in toks if w in t) * 3 + (2 if MARK.search(t) else 0)
        if a.get("조번호") == str(exclude_no): continue
        if s: scored.append((s, a))
    scored.sort(key=lambda x: -x[0])
    out = []
    for s, a in scored[:k]:
        t = art_text(a)
        m = MARK.search(t)
        snip = t[max(0, (m.start() if m else 0) - 60):(m.start() if m else 0) + 120]
        out.append(f"제{a['조번호']}조{('(' + a['조제목'] + ')') if a.get('조제목') else ''} … {snip}")
    return "\n".join(out)


def make_blank():
    ext, d4, laws, rules = load()
    rows = []
    for i, r in enumerate(sorted(d4, key=lambda x: x["행정규칙명"]), 1):
        sid = r["행정규칙일련번호"]; tri = ext[sid]; rule = rules.get(sid, {})
        ln = norm(r["주장법령"]); law = laws.get(ln)
        if law is None and ln == "학점인정법": law = laws.get(norm("학점인정 등에 관한 법률"))
        art = (r["주장조문"] or "").strip()
        status, txt, cand = "", "", ""
        if law is None:
            status = "법령 코퍼스 밖" if ln != norm(r["행정규칙명"]) else "자기 자신을 근거로 파싱됨(주장 없음과 같음)"
        elif art:
            a = find_article(law, art)
            if a and art_text(a):
                txt = art_text(a); status = "삭제된 조문" if re.match(r"^제\d+조\s*삭제", txt) else "조문 확인"
            else:
                status = f"제{art}조 현행 텍스트 없음(삭제·개정?)"
        else:
            status = "조문 번호 없이 법령만 주장"
        cits = parse_citations(first_article(rule), laws)
        pool = cited_laws(first_article(rule), laws) + ([law] if law else [])
        seen_l, kw = set(), []
        for lw in pool:
            if lw["법령명한글"] in seen_l: continue
            seen_l.add(lw["법령명한글"])
            c = candidates(lw, r["행정규칙명"], k=3)
            if c: kw.append(f"《{lw['법령명한글']}》\n{c}")
        cand = "\n".join(kw)
        cit_txt = "\n\n".join(f"[{lab}] ({st}) {tx[:1200]}" for lab, st, tx in cits) if cits else "(제1조에서 「」 인용 없음)"
        cit_mark = ",".join(sorted({m for _, st, tx in cits if st == "확인" for m in MARK.findall(tx)}))
        rows.append({
            "id": f"C{i:02d}", "행정규칙일련번호": sid, "행정규칙명": r["행정규칙명"], "종류": r["종류"],
            "발령일자": tri.get("발령일자", ""), "트리아지_T2메모": tri.get("T2_메모", ""),
            "규칙_제1조_전문": first_article(rule),
            "규칙이_인용한_조문_전문(제1조 파싱)": cit_txt, "인용조문_위임어_자동": cit_mark,
            "D4_주장법령(자동)": r["주장법령"], "D4_주장조문(자동)": art or "", "법령명_코퍼스": law["법령명한글"] if law else "",
            "D4_상태_자동": status, "D4_주장조문_현행텍스트": txt,
            "D4_자동_위임어": ",".join(sorted(set(MARK.findall(txt)))) if txt else "",
            "후보조문_자동(참고)": cand,
            "판정": "", "하위코드": "", "근거조문(저자기재)": "", "메모(저자)": "",
        })
    with open(BLANK, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    print(f"[산출] {BLANK}  {len(rows)}건")
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "claim_check"
        keys = list(rows[0].keys()); ws.append(keys)
        for r in rows: ws.append([r[k] for k in keys])
        widths = {"id": 6, "행정규칙명": 34, "종류": 6, "발령일자": 10, "트리아지_T2메모": 30, "규칙_제1조_전문": 50,
                  "규칙이_인용한_조문_전문(제1조 파싱)": 80, "인용조문_위임어_자동": 16,
                  "D4_주장법령(자동)": 22, "D4_주장조문(자동)": 8, "법령명_코퍼스": 22, "D4_상태_자동": 18, "D4_주장조문_현행텍스트": 60,
                  "D4_자동_위임어": 14, "후보조문_자동(참고)": 60, "판정": 8, "하위코드": 14, "근거조문(저자기재)": 14, "메모(저자)": 40}
        for j, k in enumerate(keys, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = widths.get(k, 14)
        yellow = PatternFill("solid", fgColor="FFFF00")
        for j, k in enumerate(keys, 1):
            c = ws.cell(row=1, column=j); c.font = Font(name="Arial", bold=True)
            if k in ("판정", "하위코드", "근거조문(저자기재)", "메모(저자)"): c.fill = yellow
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"a,b,c"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"b-decree,b-authorise,b-unrelated,b-deleted,b-cite"', allow_blank=True)
        ws.add_data_validation(dv); ws.add_data_validation(dv2)
        jp = keys.index("판정") + 1; js = keys.index("하위코드") + 1
        L = openpyxl.utils.get_column_letter
        dv.add(f"{L(jp)}2:{L(jp)}{len(rows)+1}"); dv2.add(f"{L(js)}2:{L(js)}{len(rows)+1}")
        ws.freeze_panes = "D2"
        lg = wb.create_sheet("README")
        for line in [
            "노란 열 네 개만 채운다: 판정(a/b/c) · 하위코드(b일 때) · 근거조문(저자기재) · 메모.",
            "1차 증거는 '규칙이_인용한_조문_전문(제1조 파싱)' — 규칙 스스로 근거로 든 조문의 현행 원문이다.",
            "D4_* 열은 자동 추출기(findings_D4)의 주장 조문으로, 가지번호를 놓치거나 정의 인용을 근거로 잡은 경우가 있어 참고만 한다.",
            "판정 기준: 코드북 v1.1 의 L1(위임 여부) 테스트를 '주장 조문'에 적용한다. 규칙의 규율 사항을",
            "장관 또는 행정규칙에 맡기는 문언이 있으면 a. '후보조문_자동'은 참고용이며, 같은 법령의 다른 조문에",
            "위임이 있으면 a 로 판정하고 근거조문에 그 조번호를 적는다.",
            "b-decree: 위임이 대통령령·부령까지만 간다 / b-authorise: 기준 마련·통보 권한만(2020두43722 유형) /",
            "b-unrelated: 주장 조문이 무관(번호 밀림 포함) / b-deleted: 삭제 / b-cite: 단순 인용·유보.",
            "c: 법령이 코퍼스 밖이거나 별표에만 있어 판단 불가.",
            "다 채우면 data/claim_check_labeled.csv 로 저장(엑셀 → CSV UTF-8) 후 python3 src/claim_check.py --analyse",
        ]: lg.append([line])
        lg.column_dimensions["A"].width = 110
        wb.save(BLANK.with_suffix(".xlsx")); print(f"[산출] {BLANK.with_suffix('.xlsx')}")
    except ImportError:
        print("openpyxl 없음 — CSV 만 생성")


def analyse():
    rows = list(csv.DictReader(open(LABELED, encoding="utf-8-sig")))
    n = len(rows); c = Counter(r["판정"].strip() for r in rows); sub = Counter(r["하위코드"].strip() for r in rows if r["판정"].strip() == "b")
    assert all(r["판정"].strip() in ("a", "b", "c") for r in rows), "판정 미기재 행이 있음"
    lines = [f"# 주장 조문 원문 대조 — 등록 외 사후 분석 ({n}건)\n",
             "대상: D1 ∩ 대외구속지향 61건 중 교육부 법령을 근거로 주장하는 규칙 전부.\n",
             "| 판정 | n | 뜻 |", "|---|---|---|",
             f"| a | {c['a']} | 주장 법령 원문에 위임 문언 있음 → 정부 메타데이터 누락 |",
             f"| b | {c['b']} | 원문에 위임 없음 → 근거 과장 ({', '.join(f'{k} {v}' for k, v in sorted(sub.items()))}) |",
             f"| c | {c['c']} | 판단 불가 |", "", "## 건별\n",
             "| id | 규칙 | 종류 | 근거조문(저자 확인) | 판정 | 하위 | 저자 메모 |", "|---|---|---|---|---|---|---|"]
    for r in rows:
        lines.append(f"| {r['id']} | {r['행정규칙명'][:40]} | {r['종류']} | {r['근거조문(저자기재)']} | {r['판정']} | {r['하위코드']} | {r['메모(저자)'].replace('|', '/')} |")
    lines += ["", "## 원고 반영 문장 재료\n",
              f"- Of the {n} externally-directed orphan rules that claim an MOE statute as their basis, "
              f"{c['a']} claim a provision that does contain delegating language (metadata omission), "
              f"{c['b']} claim a provision that does not (basis overstated), {c['c']} could not be decided.",
              "- 편차 로그: D-2026-09-0x — post-registration clause-level check, 등록 외 사후 분석으로 표기."]
    DOC.write_text("\n".join(lines) + "\n", encoding="utf-8"); print(f"[산출] {DOC}"); print("\n".join(lines[:8]))


if __name__ == "__main__":
    analyse() if "--analyse" in sys.argv else make_blank()
